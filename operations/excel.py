from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO

from django.contrib.auth.models import User
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Customer, Product, UserProfile

HEADER_FILL = 'E71922'


def styled_workbook(title, headers):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = max(14, len(header) + 3)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{ws.cell(1, len(headers)).coordinate}'
    return wb, ws


def workbook_bytes(wb):
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def product_template():
    headers = ['Barkod*', 'Ürün Adı*', 'Marka', 'Hacim', 'Aktif (E/H)']
    wb, ws = styled_workbook('Ürünler', headers)
    ws.append(['8690000000000', 'Örnek Ürün 100 cl', 'Örnek Marka', '100 cl', 'E'])
    return workbook_bytes(wb)


def customer_template():
    headers = [
        'Müşteri Kodu (yalnızca güncelleme)',
        'Ticari Unvan*',
        'Yetkili',
        'Telefon',
        'E-posta',
        'Satış Temsilcisi Kullanıcı Adı',
        'Notlar',
        'Aktif (E/H)',
    ]
    wb, ws = styled_workbook('Müşteriler', headers)
    ws.append(['', 'Örnek Market', 'Yetkili', '0533...', '', 'demo_satis', '', 'E'])
    ws.append(['MUS-00001', 'Mevcut müşteriyi güncelleme örneği', '', '', '', '', '', 'E'])
    return workbook_bytes(wb)


def _normalized_header(value):
    return (
        str(value or '')
        .strip()
        .lower()
        .replace('*', '')
        .replace('ı', 'i')
        .replace('ş', 's')
        .replace('ğ', 'g')
        .replace('ü', 'u')
        .replace('ö', 'o')
        .replace('ç', 'c')
    )


def _header_map(ws):
    return {_normalized_header(cell.value): idx for idx, cell in enumerate(ws[1])}


def _cell(row, mapping, *names, default=''):
    for name in names:
        idx = mapping.get(_normalized_header(name))
        if idx is not None and idx < len(row):
            value = row[idx]
            return default if value is None else value
    return default


def _is_active(value):
    return str(value or 'E').strip().upper() not in {'H', 'HAYIR', '0', 'FALSE', 'PASİF', 'PASIF'}


def import_products(file):
    wb = load_workbook(file, data_only=True)
    ws = wb.active
    mapping = _header_map(ws)
    created = updated = 0
    errors = []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not any(row):
            continue
        try:
            barcode = str(_cell(row, mapping, 'Barkod')).strip()
            name = str(_cell(row, mapping, 'Ürün Adı', 'Ürün')).strip()
            if not barcode or not name:
                raise ValueError('Barkod ve ürün adı zorunludur.')
            defaults = {
                'name': name,
                'brand': str(_cell(row, mapping, 'Marka')).strip(),
                'volume': str(_cell(row, mapping, 'Hacim')).strip(),
                'category': '',
                'units_per_case': 1,
                'base_unit': Product.Unit.PIECE,
                'is_active': _is_active(_cell(row, mapping, 'Aktif (E/H)', 'Aktif', default='E')),
            }
            _, is_created = Product.objects.update_or_create(barcode=barcode, defaults=defaults)
            created += int(is_created)
            updated += int(not is_created)
        except Exception as exc:
            errors.append(f'Satır {row_no}: {exc}')
    return created, updated, errors


def import_customers(file):
    """Yeni müşterilerde kodu sistem üretir; kod yalnızca mevcut kaydı bulmak için kullanılır."""
    wb = load_workbook(file, data_only=True)
    ws = wb.active
    mapping = _header_map(ws)
    created = updated = 0
    errors = []

    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not any(row):
            continue
        try:
            code = str(_cell(
                row,
                mapping,
                'Müşteri Kodu (yalnızca güncelleme)',
                'Müşteri Kodu',
                'Kod',
            )).strip()
            name = str(_cell(row, mapping, 'Ticari Unvan', 'Unvan')).strip()
            if not name:
                raise ValueError('Ticari unvan zorunludur.')

            salesperson_username = str(_cell(
                row,
                mapping,
                'Satış Temsilcisi Kullanıcı Adı',
                'Satış Temsilcisi',
            )).strip()
            salesperson = None
            if salesperson_username:
                salesperson = User.objects.filter(
                    username=salesperson_username,
                    is_active=True,
                    profile__active=True,
                    profile__role=UserProfile.Role.SALES,
                ).first()
                if not salesperson:
                    raise ValueError(
                        f'"{salesperson_username}" aktif bir satış temsilcisi değildir.'
                    )

            defaults = {
                'name': name,
                'contact_name': str(_cell(row, mapping, 'Yetkili', 'Yetkili Kişi')).strip(),
                'phone': str(_cell(row, mapping, 'Telefon')).strip(),
                'email': str(_cell(row, mapping, 'E-posta', 'Eposta')).strip(),
                'assigned_salesperson': salesperson,
                'notes': str(_cell(row, mapping, 'Notlar', 'Not')).strip(),
                'is_active': _is_active(_cell(row, mapping, 'Aktif (E/H)', 'Aktif', default='E')),
            }

            customer = Customer.objects.filter(code=code).first() if code else None
            if code and customer is None:
                raise ValueError(
                    f'{code} kodlu müşteri bulunamadı. Yeni müşteri için kod hücresini boş bırakın.'
                )
            if customer:
                for field, value in defaults.items():
                    setattr(customer, field, value)
                customer.save()
                updated += 1
            else:
                Customer.objects.create(**defaults)
                created += 1
        except Exception as exc:
            errors.append(f'Satır {row_no}: {exc}')
    return created, updated, errors


def export_products():
    wb, ws = styled_workbook('Ürünler', ['Barkod', 'Ürün Adı', 'Marka', 'Hacim', 'Aktif'])
    for product in Product.objects.all():
        ws.append([
            product.barcode,
            product.name,
            product.brand,
            product.volume,
            'E' if product.is_active else 'H',
        ])
    return workbook_bytes(wb)


def export_customers(queryset=None):
    wb, ws = styled_workbook(
        'Müşteriler',
        [
            'Müşteri Kodu',
            'Ticari Unvan',
            'Yetkili',
            'Telefon',
            'E-posta',
            'Satış Temsilcisi Kullanıcı Adı',
            'Notlar',
            'Aktif',
        ],
    )
    for customer in (queryset if queryset is not None else Customer.objects.all()):
        ws.append([
            customer.code,
            customer.name,
            customer.contact_name,
            customer.phone,
            customer.email,
            customer.assigned_salesperson.username if customer.assigned_salesperson else '',
            customer.notes,
            'E' if customer.is_active else 'H',
        ])
    return workbook_bytes(wb)


def export_balances(queryset=None):
    from .models import ConnectionItem

    wb, ws = styled_workbook(
        'Bağlantı Bakiyeleri',
        [
            'Müşteri Kodu',
            'Müşteri',
            'Fatura',
            'Fatura Tarihi',
            'Barkod',
            'Ürün',
            'Toplam Hak',
            'Birim Fiyat',
            'İskonto (%)',
            'KDV (%)',
            'Satır Toplamı (KDV Dahil)',
            'Teslim',
            'Sipariş',
            'Kullanılabilir',
            'Birim',
        ],
    )
    lines = queryset if queryset is not None else ConnectionItem.objects.select_related(
        'connection__customer', 'product'
    ).exclude(connection__status='CANCELLED')
    for item in lines:
        ws.append([
            item.connection.customer.code,
            item.connection.customer.name,
            item.connection.invoice_no,
            item.connection.invoice_date,
            item.product.barcode,
            item.product.name,
            int(item.effective_quantity),
            item.unit_price or 0,
            item.discount_rate,
            item.vat_rate,
            item.total_amount,
            int(item.delivered_quantity),
            int(item.reserved_quantity),
            int(item.available_quantity),
            'Adet',
        ])
    return workbook_bytes(wb)


QTY_FORMAT = '#,##0'
MONEY_FORMAT = '#,##0.00'
DATE_FORMAT = 'DD.MM.YYYY'
TOTAL_FILL = 'F4F5F7'
LINK_FONT = Font(color='0563C1', underline='single')
INVALID_SHEET_CHARS = set('[]:*?/\\')


def _sheet_title(text, used):
    clean = ''.join('-' if ch in INVALID_SHEET_CHARS else ch for ch in text).strip()[:31] or 'Müşteri'
    title, n = clean, 2
    while title.lower() in used:
        suffix = f' ({n})'
        title = clean[:31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def _write_header(ws, row, headers):
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row, col, header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _write_row(ws, row, values, formats=None, bold=False, fill=None):
    for col, value in enumerate(values, 1):
        cell = ws.cell(row, col, value)
        fmt = (formats or {}).get(col)
        if fmt:
            cell.number_format = fmt
        if bold:
            cell.font = Font(bold=True)
        if fill:
            cell.fill = PatternFill('solid', fgColor=fill)


def _set_widths(ws, widths):
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = width


LINE_TOTAL_KEYS = ['effective', 'delivered', 'reserved', 'remaining', 'available', 'total_amount', 'remaining_amount']


def _line_figures(item):
    remaining = int(item.remaining_quantity)
    unit_total = (item.total_amount / item.quantity) if item.quantity else Decimal('0')
    return {
        'effective': int(item.effective_quantity),
        'delivered': int(item.delivered_quantity),
        'reserved': int(item.reserved_quantity),
        'remaining': remaining,
        'available': int(item.available_quantity),
        'total_amount': item.total_amount,
        'remaining_amount': (unit_total * remaining).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP),
    }


def export_connection_book(customers):
    """Tek Excel dosyası: Özet sekmesi + her müşteri için bağlantı satırları ve teslim geçmişi sekmesi."""
    from .models import ConnectionItem, ConnectionSale, Shipment, ShipmentItem

    wb = Workbook()
    summary = wb.active
    summary.title = 'Özet'
    used_titles = {'özet'}
    today = timezone.localdate()

    summary.cell(1, 1, 'Bağlantı Raporu').font = Font(bold=True, size=14)
    summary.cell(2, 1, f'Rapor tarihi: {today:%d.%m.%Y}')
    _write_header(summary, 4, [
        'Müşteri Kodu', 'Müşteri', 'Bağlantı Sayısı', 'Toplam Hak', 'Teslim Edilen',
        'Siparişte Bekleyen', 'Kalan', 'Siparişe Açık', 'Toplam Tutar (KDV Dahil)', 'Kalan Tutar (KDV Dahil)', 'Sekme',
    ])
    summary_formats = {3: QTY_FORMAT, 4: QTY_FORMAT, 5: QTY_FORMAT, 6: QTY_FORMAT, 7: QTY_FORMAT, 8: QTY_FORMAT,
                       9: MONEY_FORMAT, 10: MONEY_FORMAT}
    summary_row = 5
    grand_connections = 0
    grand = dict.fromkeys(LINE_TOTAL_KEYS, 0)

    line_headers = [
        'Fatura No', 'Fatura Tarihi', 'Durum', 'Barkod', 'Ürün', 'Bağlantı Miktarı', 'Teslim Edilen',
        'Siparişte Bekleyen', 'Kalan', 'Siparişe Açık', 'Birim Fiyat', 'İskonto (%)', 'KDV (%)',
        'Satır Toplamı (KDV Dahil)', 'Kalan Tutar (KDV Dahil)',
    ]
    line_formats = {2: DATE_FORMAT, 6: QTY_FORMAT, 7: QTY_FORMAT, 8: QTY_FORMAT, 9: QTY_FORMAT, 10: QTY_FORMAT,
                    11: MONEY_FORMAT, 14: MONEY_FORMAT, 15: MONEY_FORMAT}
    history_headers = [
        'Teslim Tarihi', 'Sevk No', 'Sipariş No', 'Fatura No', 'Barkod', 'Ürün', 'Teslim Miktarı',
        'Teslim Alan', 'Teslim Eden',
    ]

    for customer in customers.order_by('name'):
        connections = list(
            ConnectionSale.objects.filter(customer=customer)
            .exclude(status=ConnectionSale.Status.CANCELLED)
            .order_by('invoice_date', 'id')
        )
        if not connections:
            continue

        ws = wb.create_sheet(_sheet_title(f'{customer.code} {customer.name}', used_titles))
        ws.cell(1, 1, customer.name).font = Font(bold=True, size=14)
        ws.cell(2, 1, f'{customer.code} • Rapor tarihi: {today:%d.%m.%Y}')
        back = ws.cell(2, 6, '← Özet sekmesine dön')
        back.hyperlink = "#'Özet'!A1"
        back.font = LINK_FONT

        ws.cell(4, 1, 'BAĞLANTI SATIRLARI').font = Font(bold=True)
        _write_header(ws, 5, line_headers)
        row = 6
        totals = dict.fromkeys(LINE_TOTAL_KEYS, 0)
        items = ConnectionItem.objects.select_related('connection', 'product').filter(
            connection__in=connections
        ).order_by('connection__invoice_date', 'connection_id', 'product__name', 'id')
        for item in items:
            f = _line_figures(item)
            _write_row(ws, row, [
                item.connection.invoice_no,
                item.connection.invoice_date,
                item.connection.get_status_display(),
                item.product.barcode,
                item.product.name,
                f['effective'], f['delivered'], f['reserved'], f['remaining'], f['available'],
                item.unit_price or 0,
                item.discount_rate,
                item.vat_rate,
                f['total_amount'],
                f['remaining_amount'],
            ], line_formats)
            for key in LINE_TOTAL_KEYS:
                totals[key] += f[key]
            row += 1
        _write_row(ws, row, [
            'TOPLAM', '', '', '', '',
            totals['effective'], totals['delivered'], totals['reserved'], totals['remaining'], totals['available'],
            '', '', '', totals['total_amount'], totals['remaining_amount'],
        ], line_formats, bold=True, fill=TOTAL_FILL)

        row += 3
        ws.cell(row, 1, 'TESLİM GEÇMİŞİ').font = Font(bold=True)
        row += 1
        _write_header(ws, row, history_headers)
        row += 1
        deliveries = ShipmentItem.objects.select_related(
            'shipment__order', 'shipment__delivered_by', 'connection_item__connection', 'connection_item__product'
        ).filter(
            shipment__status=Shipment.Status.DELIVERED,
            connection_item__connection__in=connections,
        ).order_by('shipment__delivery_date', 'shipment__number', 'id')
        history_start = row
        for line in deliveries:
            shipment = line.shipment
            deliverer = shipment.delivered_by
            _write_row(ws, row, [
                shipment.delivery_date,
                shipment.number,
                shipment.order.number,
                line.connection_item.connection.invoice_no,
                line.connection_item.product.barcode,
                line.connection_item.product.name,
                int(line.quantity_delivered),
                shipment.delivered_to,
                (deliverer.get_full_name() or deliverer.username) if deliverer else '',
            ], {1: DATE_FORMAT, 7: QTY_FORMAT})
            row += 1
        if row == history_start:
            ws.cell(row, 1, 'Henüz teslimat yok.')

        _set_widths(ws, [16, 13, 20, 16, 38, 13, 13, 13, 11, 13, 12, 11, 9, 18, 18])
        ws.freeze_panes = 'A6'

        _write_row(summary, summary_row, [
            customer.code, customer.name, len(connections),
            totals['effective'], totals['delivered'], totals['reserved'], totals['remaining'], totals['available'],
            totals['total_amount'], totals['remaining_amount'], ws.title,
        ], summary_formats)
        link = summary.cell(summary_row, 11)
        link.hyperlink = f"#'{ws.title}'!A1"
        link.font = LINK_FONT
        grand_connections += len(connections)
        for key in LINE_TOTAL_KEYS:
            grand[key] += totals[key]
        summary_row += 1

    if summary_row == 5:
        summary.cell(5, 1, 'Raporlanacak bağlantı bulunamadı.')
    else:
        _write_row(summary, summary_row, [
            'TOPLAM', '', grand_connections,
            grand['effective'], grand['delivered'], grand['reserved'], grand['remaining'], grand['available'],
            grand['total_amount'], grand['remaining_amount'], '',
        ], summary_formats, bold=True, fill=TOTAL_FILL)
    _set_widths(summary, [14, 36, 12, 13, 13, 13, 11, 13, 20, 20, 30])
    summary.freeze_panes = 'A5'
    return workbook_bytes(wb)
