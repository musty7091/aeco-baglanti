from decimal import Decimal
from io import BytesIO
import tempfile

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone

from .forms import (
    ConnectionItemForm,
    ConnectionSaleForm,
    CustomerForm,
    ShipmentDeliveryForm,
    ShipmentForm,
)
from .models import (
    BalanceAdjustment,
    ConnectionItem,
    ConnectionSale,
    Customer,
    DocumentCounter,
    Order,
    Product,
    Shipment,
    ShipmentItem,
    UserProfile,
    Vehicle,
    Warehouse,
)
from openpyxl import load_workbook
from PIL import Image

from .pdf import shipment_pdf
from .services import create_order_from_request, deliver_shipment, parse_quantity, shipment_balance_rows


class CustomerCodeAndFormTests(TestCase):
    def setUp(self):
        DocumentCounter.objects.update_or_create(
            key='customer', defaults={'year': 0, 'value': 0}
        )

    def test_customer_codes_are_automatic_and_sequential(self):
        first = Customer.objects.create(name='Birinci Müşteri')
        second = Customer.objects.create(name='İkinci Müşteri')
        third = Customer.objects.create(name='Üçüncü Müşteri')

        self.assertEqual(first.code, 'MUS-00001')
        self.assertEqual(second.code, 'MUS-00002')
        self.assertEqual(third.code, 'MUS-00003')
        self.assertEqual(DocumentCounter.objects.get(key='customer').value, 3)

    def test_customer_form_has_only_requested_business_fields(self):
        form = CustomerForm()
        self.assertEqual(
            list(form.fields),
            [
                'name', 'contact_name', 'phone', 'email',
                'assigned_salesperson', 'notes', 'is_active',
            ],
        )
        for removed in (
            'code', 'billing_address', 'delivery_addresses',
            'delivery_addresses_text',
        ):
            self.assertNotIn(removed, form.fields)

    def test_customer_model_has_no_removed_address_fields(self):
        field_names = {field.name for field in Customer._meta.get_fields()}
        self.assertNotIn('billing_address', field_names)
        self.assertNotIn('delivery_addresses', field_names)

    def test_customer_code_cannot_be_changed_through_edit_form(self):
        customer = Customer.objects.create(name='Kod Koruma Testi')
        original_code = customer.code
        form = CustomerForm(
            data={
                'name': 'Kod Koruma Testi Güncel',
                'contact_name': '',
                'phone': '',
                'email': '',
                'assigned_salesperson': '',
                'notes': '',
                'is_active': 'on',
                'code': 'HACK-99999',
            },
            instance=customer,
        )
        self.assertTrue(form.is_valid(), form.errors)
        updated = form.save()
        self.assertEqual(updated.code, original_code)


class InvoiceAmountAndAttachmentTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            'invoice_admin', 'invoice@example.com', 'StrongPass123!'
        )
        self.customer = Customer.objects.create(name='Fatura Test Müşterisi')
        self.product_a = Product.objects.create(
            barcode='8690000000001', name='Test Ürün A'
        )
        self.product_b = Product.objects.create(
            barcode='8690000000002', name='Test Ürün B'
        )
        self.connection = ConnectionSale.objects.create(
            number='BGL-TEST-1',
            customer=self.customer,
            invoice_no='FAT-TEST-1',
            invoice_date=timezone.localdate(),
            status=ConnectionSale.Status.PENDING,
            created_by=self.user,
        )

    def create_item(
        self, product, quantity, unit_price, discount_rate=0, vat_rate=20
    ):
        return ConnectionItem.objects.create(
            connection=self.connection,
            product=product,
            quantity=quantity,
            unit_price=Decimal(str(unit_price)),
            discount_rate=Decimal(str(discount_rate)),
            vat_rate=Decimal(str(vat_rate)),
        )

    def test_connection_form_has_no_expiry_date(self):
        self.assertNotIn('expiry_date', ConnectionSaleForm().fields)
        field_names = {field.name for field in ConnectionSale._meta.get_fields()}
        self.assertNotIn('expiry_date', field_names)

    def test_line_total_calculates_discount_and_vat(self):
        line = self.create_item(
            self.product_a, quantity=10, unit_price='100.00',
            discount_rate='10.00', vat_rate='20.00',
        )
        self.assertEqual(line.gross_amount, Decimal('1000.00'))
        self.assertEqual(line.discount_amount, Decimal('100.00'))
        self.assertEqual(line.net_amount, Decimal('900.00'))
        self.assertEqual(line.vat_amount, Decimal('180.00'))
        self.assertEqual(line.total_amount, Decimal('1080.00'))

    def test_connection_totals_sum_multiple_lines(self):
        self.create_item(
            self.product_a, quantity=2, unit_price='100.00',
            discount_rate='10.00', vat_rate='20.00',
        )
        self.create_item(
            self.product_b, quantity=3, unit_price='50.00',
            discount_rate='0.00', vat_rate='10.00',
        )
        self.assertEqual(self.connection.gross_amount, Decimal('350.00'))
        self.assertEqual(self.connection.discount_amount, Decimal('20.00'))
        self.assertEqual(self.connection.net_amount, Decimal('330.00'))
        self.assertEqual(self.connection.vat_amount, Decimal('51.00'))
        self.assertEqual(self.connection.total_amount, Decimal('381.00'))

    def test_fully_discounted_free_line_does_not_increase_total(self):
        paid = self.create_item(
            self.product_a, quantity=1, unit_price='100.00',
            discount_rate='0.00', vat_rate='20.00',
        )
        free = self.create_item(
            self.product_b, quantity=50, unit_price='999.00',
            discount_rate='100.00', vat_rate='20.00',
        )
        self.assertEqual(paid.total_amount, Decimal('120.00'))
        self.assertEqual(free.total_amount, Decimal('0.00'))
        self.assertTrue(free.is_free_item)
        self.assertEqual(self.connection.total_amount, Decimal('120.00'))

    def _form_data(self):
        return {
            'customer': self.customer.pk,
            'invoice_no': 'FAT-DOSYA-1',
            'invoice_date': timezone.localdate().isoformat(),
            'notes': '',
        }

    def test_invoice_attachment_extension_validation(self):
        bad_file = SimpleUploadedFile(
            'fatura.exe', b'not-a-document', content_type='application/octet-stream'
        )
        form = ConnectionSaleForm(
            data=self._form_data(), files={'invoice_attachment': bad_file}
        )
        self.assertFalse(form.is_valid())
        self.assertIn('invoice_attachment', form.errors)

    def test_invoice_attachment_size_validation(self):
        oversized = SimpleUploadedFile(
            'fatura.pdf', b'x' * (10 * 1024 * 1024 + 1),
            content_type='application/pdf',
        )
        form = ConnectionSaleForm(
            data=self._form_data(), files={'invoice_attachment': oversized}
        )
        self.assertFalse(form.is_valid())
        self.assertIn('invoice_attachment', form.errors)

    def test_supported_invoice_attachment_is_valid(self):
        valid_pdf = SimpleUploadedFile(
            'fatura.pdf', b'%PDF-1.4\n%%EOF', content_type='application/pdf'
        )
        form = ConnectionSaleForm(
            data=self._form_data(), files={'invoice_attachment': valid_pdf}
        )
        self.assertTrue(form.is_valid(), form.errors)


class CreateViewTests(TestCase):
    def setUp(self):
        DocumentCounter.objects.update_or_create(
            key='customer', defaults={'year': 0, 'value': 0}
        )
        self.user = User.objects.create_superuser(
            'create_admin', 'create@example.com', 'StrongPass123!'
        )
        self.client.force_login(self.user)

    def test_new_customer_creation_works_and_assigns_code(self):
        response = self.client.post(
            reverse('customer_create'),
            {
                'name': 'Yeni Ticari Müşteri Ltd.',
                'contact_name': 'Ayşe Şahin',
                'phone': '0533 123 45 67',
                'email': 'ayse@example.com',
                'assigned_salesperson': '',
                'notes': 'Yeni müşteri testi',
                'is_active': 'on',
            },
        )
        self.assertEqual(response.status_code, 302)
        customer = Customer.objects.get(name='Yeni Ticari Müşteri Ltd.')
        self.assertEqual(customer.code, 'MUS-00001')

    def test_new_connection_sale_creation_works(self):
        customer = Customer.objects.create(name='Bağlantı Müşterisi')
        product = Product.objects.create(
            barcode='9900000000999', name='Bağlantı Test Ürünü'
        )
        response = self.client.post(
            reverse('connection_create'),
            {
                'customer': customer.pk,
                'invoice_no': 'TEST-FTR-2026-1',
                'invoice_date': timezone.localdate().isoformat(),
                'notes': 'Bağlantı oluşturma testi',
                'items-TOTAL_FORMS': '1',
                'items-INITIAL_FORMS': '0',
                'items-MIN_NUM_FORMS': '1',
                'items-MAX_NUM_FORMS': '1000',
                'items-0-product': str(product.pk),
                'items-0-quantity': '4',
                'items-0-unit_price': '125.00',
                'items-0-discount_rate': '10.00',
                'items-0-vat_rate': '20.00',
            },
        )
        self.assertEqual(response.status_code, 302)
        connection = ConnectionSale.objects.get(invoice_no='TEST-FTR-2026-1')
        self.assertEqual(connection.items.count(), 1)
        self.assertEqual(connection.total_amount, Decimal('540.00'))


class BusinessFlowTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            'admin', 'a@a.com', 'StrongPass123!'
        )
        self.customer = Customer.objects.create(name='Test Market')
        self.product = Product.objects.create(
            barcode='8690000000101', name='Test Ürün'
        )
        self.warehouse = Warehouse.objects.create(code='ANA', name='Ana Depo')
        self.connection = ConnectionSale.objects.create(
            number='BGL-1', customer=self.customer, invoice_no='FAT-100',
            invoice_date=timezone.localdate(),
            status=ConnectionSale.Status.ACTIVE, created_by=self.user,
            approved_by=self.user, approved_at=timezone.now(),
        )
        self.line = ConnectionItem.objects.create(
            connection=self.connection, product=self.product,
            quantity=100, unit=Product.Unit.PIECE,
            unit_price=Decimal('100.00'), vat_rate=Decimal('20.00'),
        )

    def create_approved_order(self, quantity='25'):
        order = create_order_from_request(
            customer=self.customer,
            user=self.user,
            order_data={
                'order_date': timezone.localdate(),
                'requested_delivery_date': None,
                'delivery_address': 'Adres',
                'notes': '',
            },
            quantities={str(self.line.pk): quantity},
        )
        order.status = Order.Status.APPROVED
        order.approved_by = self.user
        order.approved_at = timezone.now()
        order.save()
        return order

    def test_connection_item_form_is_simplified(self):
        line_form = ConnectionItemForm()
        self.assertEqual(
            list(line_form.fields),
            ['product', 'quantity', 'unit_price', 'discount_rate', 'vat_rate', 'total_amount'],
        )

    def test_invoice_based_reservation_and_delivery(self):
        order = self.create_approved_order('25')
        self.assertEqual(self.line.available_quantity, 75)
        shipment = Shipment.objects.create(
            number='SVK-1', order=order, customer=self.customer,
            warehouse=self.warehouse, created_by=self.user,
        )
        order_item = order.items.get()
        ShipmentItem.objects.create(
            shipment=shipment,
            order_item=order_item,
            connection_item=self.line,
            quantity_delivered=25,
        )
        deliver_shipment(shipment=shipment, user=self.user)
        self.line.refresh_from_db()
        self.assertEqual(self.line.delivered_quantity, 25)
        self.assertEqual(self.line.available_quantity, 75)

    def test_open_shipment_allocation_prevents_duplicate_planning(self):
        order = self.create_approved_order('25')
        order_item = order.items.get()
        shipment = Shipment.objects.create(
            number='SVK-2', order=order, customer=self.customer,
            warehouse=self.warehouse, created_by=self.user,
        )
        ShipmentItem.objects.create(
            shipment=shipment,
            order_item=order_item,
            connection_item=self.line,
            quantity_delivered=10,
        )
        self.assertEqual(order_item.allocatable_quantity, 15)

    def test_adjustment_cannot_reduce_below_reserved_amount(self):
        self.create_approved_order('25')
        adjustment = BalanceAdjustment(
            connection_item=self.line,
            signed_quantity=-80,
            kind=BalanceAdjustment.Kind.CORRECTION,
            reason='Test',
            created_by=self.user,
            approved_by=self.user,
        )
        with self.assertRaises(ValidationError):
            adjustment.full_clean()

    def test_quantities_are_whole_numbers(self):
        self.assertEqual(parse_quantity('18'), 18)
        self.assertEqual(parse_quantity('18,000'), 18)
        with self.assertRaises(ValidationError):
            parse_quantity('18,5')

    def test_shipment_form_has_only_date(self):
        shipment = Shipment(
            order=self.create_approved_order('5'),
            customer=self.customer,
            warehouse=self.warehouse,
            created_by=self.user,
        )
        form = ShipmentForm(instance=shipment)
        self.assertEqual(list(form.fields), ['warehouse', 'shipment_date'])
        # Tek aktif depo varsa seçim gizlenir.
        self.assertTrue(form.fields['warehouse'].widget.is_hidden)

    def test_delivery_form_photos_optional_and_notes_removed(self):
        shipment = Shipment(
            order=self.create_approved_order('5'),
            customer=self.customer,
            warehouse=self.warehouse,
            created_by=self.user,
        )
        form = ShipmentDeliveryForm(instance=shipment)
        self.assertEqual(list(form.fields), ['delivery_date', 'delivered_to', 'delivery_proof'])
        self.assertFalse(form.fields['delivery_proof'].required)


class DemoDataCommandTests(TestCase):
    def setUp(self):
        self.temp_media = tempfile.TemporaryDirectory()
        self.media_override = self.settings(MEDIA_ROOT=self.temp_media.name)
        self.media_override.enable()

    def tearDown(self):
        self.media_override.disable()
        self.temp_media.cleanup()

    def test_demo_command_runs_with_new_model(self):
        call_command('demo_aeco', '--reset', verbosity=0)

        customers = Customer.objects.filter(notes__startswith='DEMO:').order_by('code')
        self.assertEqual(customers.count(), 3)
        self.assertEqual(
            list(customers.values_list('code', flat=True)),
            ['MUS-00001', 'MUS-00002', 'MUS-00003'],
        )
        self.assertEqual(Product.objects.filter(barcode__startswith='990000').count(), 10)
        self.assertTrue(Product.objects.filter(name__contains='Jägermeister').exists())
        self.assertTrue(Product.objects.filter(name__contains='Moët & Chandon').exists())
        self.assertEqual(ConnectionSale.objects.filter(number__startswith='DEMO-').count(), 3)
        self.assertEqual(Order.objects.filter(number__startswith='DEMO-').count(), 4)
        self.assertEqual(Shipment.objects.filter(number__startswith='DEMO-').count(), 2)
        self.assertEqual(Vehicle.objects.filter(plate__startswith='DEMO-').count(), 2)
        self.assertEqual(
            User.objects.get(username='demo_sevkiyat').last_name,
            'Şahin',
        )

        pit_connection = ConnectionSale.objects.get(number='DEMO-BGL-0001')
        repeated_barcode = Product.objects.get(barcode='9900000000001')
        repeated_lines = ConnectionItem.objects.filter(
            connection=pit_connection, product=repeated_barcode
        )
        self.assertEqual(repeated_lines.count(), 2)
        self.assertTrue(repeated_lines.filter(discount_rate=100).exists())
        self.assertGreater(pit_connection.total_amount, Decimal('0.00'))

        call_command('demo_aeco', '--clear', verbosity=0)
        self.assertFalse(Customer.objects.filter(notes__startswith='DEMO:').exists())
        self.assertFalse(User.objects.filter(username='demo_admin').exists())


class RolePermissionTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user('role_admin', password='Pass12345!')
        self.admin.profile.role = UserProfile.Role.ADMIN
        self.admin.profile.save()

        self.management = User.objects.create_user(
            'role_management', password='Pass12345!'
        )
        self.management.profile.role = UserProfile.Role.MANAGEMENT
        self.management.profile.save()

        self.sales = User.objects.create_user(
            'role_sales', password='Pass12345!', first_name='Satış'
        )
        self.sales.profile.role = UserProfile.Role.SALES
        self.sales.profile.save()

        self.other_sales = User.objects.create_user(
            'role_sales_2', password='Pass12345!'
        )
        self.other_sales.profile.role = UserProfile.Role.SALES
        self.other_sales.profile.save()

        self.warehouse_user = User.objects.create_user(
            'role_warehouse', password='Pass12345!'
        )
        self.warehouse_user.profile.role = UserProfile.Role.WAREHOUSE
        self.warehouse_user.profile.save()

        self.dispatch = User.objects.create_user(
            'role_dispatch', password='Pass12345!', first_name='Sevk'
        )
        self.dispatch.profile.role = UserProfile.Role.DISPATCH
        self.dispatch.profile.save()

        self.other_dispatch = User.objects.create_user(
            'role_dispatch_2', password='Pass12345!'
        )
        self.other_dispatch.profile.role = UserProfile.Role.DISPATCH
        self.other_dispatch.profile.save()

        self.customer = Customer.objects.create(
            name='Satış Müşterisi', assigned_salesperson=self.sales
        )
        self.other_customer = Customer.objects.create(
            name='Diğer Müşteri', assigned_salesperson=self.other_sales
        )
        self.unassigned_customer = Customer.objects.create(name='Atanmamış Müşteri')
        self.product = Product.objects.create(
            barcode='9900990099001', name='Yetki Test Ürünü'
        )
        self.warehouse = Warehouse.objects.create(
            code='ROLE-WH', name='Yetki Test Deposu'
        )

        self.connection = ConnectionSale.objects.create(
            number='ROLE-BGL-1',
            customer=self.customer,
            invoice_no='ROLE-FAT-1',
            invoice_date=timezone.localdate(),
            status=ConnectionSale.Status.ACTIVE,
            created_by=self.management,
            approved_by=self.management,
            approved_at=timezone.now(),
        )
        self.connection_item = ConnectionItem.objects.create(
            connection=self.connection,
            product=self.product,
            quantity=50,
            unit_price=100,
            vat_rate=20,
        )
        self.pending_order = Order.objects.create(
            number='ROLE-SIP-P',
            customer=self.customer,
            status=Order.Status.PENDING,
            created_by=self.sales,
        )
        self.approved_order = Order.objects.create(
            number='ROLE-SIP-A',
            customer=self.customer,
            status=Order.Status.APPROVED,
            created_by=self.sales,
            approved_by=self.management,
            approved_at=timezone.now(),
        )
        self.order_item = self.approved_order.items.create(
            connection_item=self.connection_item,
            quantity_requested=10,
        )
        self.assigned_shipment = Shipment.objects.create(
            number='ROLE-SVK-1',
            order=self.approved_order,
            customer=self.customer,
            warehouse=self.warehouse,
            created_by=self.warehouse_user,
            assigned_dispatcher=self.dispatch,
            status=Shipment.Status.LOADED,
        )
        ShipmentItem.objects.create(
            shipment=self.assigned_shipment,
            order_item=self.order_item,
            connection_item=self.connection_item,
            quantity_delivered=10,
        )
        self.other_shipment = Shipment.objects.create(
            number='ROLE-SVK-2',
            order=self.approved_order,
            customer=self.customer,
            warehouse=self.warehouse,
            created_by=self.warehouse_user,
            assigned_dispatcher=self.other_dispatch,
            status=Shipment.Status.LOADED,
        )

    def test_sales_sees_only_explicitly_assigned_customers(self):
        self.client.force_login(self.sales)
        response = self.client.get(reverse('customer_list'))
        self.assertContains(response, 'Satış Müşterisi')
        self.assertNotContains(response, 'Diğer Müşteri')
        self.assertNotContains(response, 'Atanmamış Müşteri')
        self.assertEqual(
            self.client.get(reverse('customer_detail', args=[self.other_customer.pk])).status_code,
            404,
        )

    def test_warehouse_sees_operational_orders_but_not_pending_orders(self):
        self.client.force_login(self.warehouse_user)
        response = self.client.get(reverse('order_list'))
        self.assertContains(response, 'ROLE-SIP-A')
        self.assertNotContains(response, 'ROLE-SIP-P')

    def test_dispatch_sees_all_shipments(self):
        # Sevkiyatlar personele atanmadığı için her sevkiyat personeli tüm sevkiyatları görür.
        self.client.force_login(self.dispatch)
        response = self.client.get(reverse('shipment_list'))
        self.assertContains(response, 'ROLE-SVK-1')
        self.assertContains(response, 'ROLE-SVK-2')


class WarehouseManagementTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user('warehouse_admin', password='Pass12345!')
        self.admin.profile.role = UserProfile.Role.ADMIN
        self.admin.profile.save()
        self.management = User.objects.create_user('warehouse_management', password='Pass12345!')
        self.management.profile.role = UserProfile.Role.MANAGEMENT
        self.management.profile.save()
        self.warehouse = Warehouse.objects.create(code='ANA', name='Ana Depo', is_active=True)

    def test_admin_can_edit_warehouse(self):
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse('warehouse_edit', args=[self.warehouse.pk]),
            {'code': 'MERKEZ', 'name': 'Merkez Depo', 'address': 'Lefkoşa', 'is_active': 'on'},
        )
        self.assertRedirects(response, reverse('warehouse_list'))
        self.warehouse.refresh_from_db()
        self.assertEqual(self.warehouse.code, 'MERKEZ')
        self.assertEqual(self.warehouse.name, 'Merkez Depo')

    def test_admin_can_toggle_warehouse(self):
        self.client.force_login(self.admin)
        response = self.client.post(reverse('warehouse_toggle', args=[self.warehouse.pk]))
        self.assertRedirects(response, reverse('warehouse_list'))
        self.warehouse.refresh_from_db()
        self.assertFalse(self.warehouse.is_active)

    def test_management_can_edit_and_toggle_warehouse(self):
        self.client.force_login(self.management)
        edit_response = self.client.post(
            reverse('warehouse_edit', args=[self.warehouse.pk]),
            {'code': 'ANA', 'name': 'Ana Depo Güncel', 'address': '', 'is_active': 'on'},
        )
        self.assertRedirects(edit_response, reverse('warehouse_list'))
        toggle_response = self.client.post(reverse('warehouse_toggle', args=[self.warehouse.pk]))
        self.assertRedirects(toggle_response, reverse('warehouse_list'))
        self.warehouse.refresh_from_db()
        self.assertEqual(self.warehouse.name, 'Ana Depo Güncel')
        self.assertFalse(self.warehouse.is_active)


class SimplifiedShipmentFlowTests(TestCase):
    def setUp(self):
        self.temp_media = tempfile.TemporaryDirectory()
        self.media_override = self.settings(MEDIA_ROOT=self.temp_media.name)
        self.media_override.enable()
        self.admin = User.objects.create_user('flow_admin', password='Pass12345!', first_name='Ayşe', last_name='Yönetici')
        self.admin.profile.role = UserProfile.Role.ADMIN
        self.admin.profile.save()
        self.driver = User.objects.create_user('flow_driver', password='Pass12345!', first_name='Mehmet')
        self.driver.profile.role = UserProfile.Role.DISPATCH
        self.driver.profile.save()
        self.warehouse = Warehouse.objects.create(code='ANA', name='Ana Depo')
        self.customer = Customer.objects.create(name='Akış Market')
        self.product_a = Product.objects.create(barcode='8690000000201', name='Ürün A')
        self.product_b = Product.objects.create(barcode='8690000000202', name='Ürün B')
        self.connection = ConnectionSale.objects.create(
            number='BGL-F1', customer=self.customer, invoice_no='FAT-F1',
            invoice_date=timezone.localdate(), status=ConnectionSale.Status.ACTIVE,
            created_by=self.admin, approved_by=self.admin, approved_at=timezone.now(),
        )
        self.line_a = ConnectionItem.objects.create(
            connection=self.connection, product=self.product_a, quantity=100,
            unit_price=Decimal('10.00'), vat_rate=Decimal('20.00'),
        )
        self.line_b = ConnectionItem.objects.create(
            connection=self.connection, product=self.product_b, quantity=50,
            unit_price=Decimal('5.00'), vat_rate=Decimal('0.00'),
        )
        self.order = create_order_from_request(
            customer=self.customer, user=self.admin,
            order_data={'order_date': timezone.localdate(), 'requested_delivery_date': None,
                        'delivery_address': 'Lefkoşa', 'notes': ''},
            quantities={str(self.line_a.pk): '30'},
        )
        self.order.status = Order.Status.APPROVED
        self.order.approved_by = self.admin
        self.order.approved_at = timezone.now()
        self.order.save()
        self.client.force_login(self.admin)

    def tearDown(self):
        self.media_override.disable()
        self.temp_media.cleanup()

    def ship(self, quantity='30'):
        order_item = self.order.items.get()
        response = self.client.post(reverse('shipment_create', args=[self.order.pk]), {
            'warehouse': self.warehouse.pk,
            'shipment_date': timezone.localdate().isoformat(),
            f'quantity_{order_item.pk}': quantity,
        })
        shipment = Shipment.objects.get(order=self.order)
        self.assertRedirects(response, reverse('shipment_detail', args=[shipment.pk]))
        return shipment

    def test_single_warehouse_is_prefilled_in_hidden_field(self):
        response = self.client.get(reverse('shipment_create', args=[self.order.pk]))
        self.assertContains(response, f'type="hidden" name="warehouse" value="{self.warehouse.pk}"')
        self.assertContains(response, f'value="{timezone.localdate().isoformat()}"')

    def test_shipment_starts_on_route_with_creator_and_order_address(self):
        shipment = self.ship()
        self.assertEqual(shipment.status, Shipment.Status.ON_ROUTE)
        self.assertEqual(shipment.prepared_by_name, 'Ayşe Yönetici')
        self.assertEqual(shipment.delivery_address, 'Lefkoşa')
        self.assertEqual(shipment.items.get().quantity_delivered, 30)

    def test_order_with_open_shipment_links_to_delivery(self):
        shipment = self.ship()
        self.order.refresh_from_db()
        self.assertEqual(self.order.get_status_display(), 'Sevkte')
        response = self.client.get(reverse('order_detail', args=[self.order.pk]))
        self.assertContains(response, reverse('shipment_deliver', args=[shipment.pk]))
        self.assertNotContains(response, reverse('shipment_create', args=[self.order.pk]))

    def test_any_driver_can_deliver_without_photo(self):
        shipment = self.ship()
        self.client.force_login(self.driver)
        response = self.client.post(reverse('shipment_deliver', args=[shipment.pk]), {
            'delivery_date': timezone.localdate().isoformat(),
            'delivered_to': 'Ali Veli',
        })
        self.assertRedirects(response, reverse('shipment_detail', args=[shipment.pk]))
        shipment.refresh_from_db()
        self.assertEqual(shipment.status, Shipment.Status.DELIVERED)
        self.assertEqual(shipment.delivered_by, self.driver)
        self.assertEqual(self.line_a.delivered_quantity, 30)

    def test_delivery_with_optional_photo(self):
        shipment = self.ship()
        image = BytesIO()
        Image.new('RGB', (40, 40), 'white').save(image, format='PNG')
        upload = SimpleUploadedFile('teslim.png', image.getvalue(), content_type='image/png')
        self.client.post(reverse('shipment_deliver', args=[shipment.pk]), {
            'delivery_date': timezone.localdate().isoformat(),
            'delivered_to': 'Ali Veli',
            'delivery_proof': upload,
        })
        shipment.refresh_from_db()
        self.assertEqual(shipment.status, Shipment.Status.DELIVERED)
        self.assertTrue(shipment.delivery_proof)

    def test_shipment_document_lists_remaining_customer_balance(self):
        shipment = self.ship('30')
        rows = {row['item'].pk: row for row in shipment_balance_rows(shipment)}
        self.assertEqual(rows[self.line_a.pk]['this'], 30)
        self.assertEqual(rows[self.line_a.pk]['remaining'], 70)
        self.assertEqual(rows[self.line_b.pk]['this'], 0)
        self.assertEqual(rows[self.line_b.pk]['remaining'], 50)
        deliver_shipment(shipment=shipment, user=self.driver)
        shipment.refresh_from_db()
        rows = {row['item'].pk: row for row in shipment_balance_rows(shipment)}
        self.assertEqual(rows[self.line_a.pk]['before'], 0)
        self.assertEqual(rows[self.line_a.pk]['remaining'], 70)
        self.assertTrue(shipment_pdf(shipment, 'http://test/verify/').startswith(b'%PDF'))

    def test_connection_book_has_summary_and_customer_tabs(self):
        shipment = self.ship('30')
        deliver_shipment(shipment=shipment, user=self.driver)
        Customer.objects.create(name='Boş Müşteri')
        response = self.client.get(reverse('data_export', args=['connection_book']))
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(wb.sheetnames[0], 'Özet')
        self.assertEqual(len(wb.sheetnames), 2)
        rows = list(wb[wb.sheetnames[1]].iter_rows(values_only=True))
        flat = [value for row in rows for value in row if value is not None]
        self.assertIn('TESLİM GEÇMİŞİ', flat)
        self.assertIn(shipment.number, flat)
        total_row = next(row for row in rows if row[0] == 'TOPLAM')
        self.assertEqual(total_row[5:10], (150, 30, 0, 120, 120))
        # Kalan tutar: A 70 x 12,00 + B 50 x 5,00 = 1090,00
        self.assertEqual(Decimal(str(total_row[14])), Decimal('1090.00'))
        summary_total = next(row for row in wb['Özet'].iter_rows(values_only=True) if row[0] == 'TOPLAM')
        self.assertEqual(summary_total[2], 1)
        self.assertEqual(summary_total[6], 120)

    def test_connection_book_single_customer_and_permissions(self):
        Customer.objects.create(name='Başka Müşteri')
        response = self.client.get(reverse('data_export', args=['connection_book']) + f'?customer={self.customer.pk}')
        self.assertEqual(len(load_workbook(BytesIO(response.content)).sheetnames), 2)
        self.client.force_login(self.driver)
        response = self.client.get(reverse('data_export', args=['connection_book']))
        self.assertRedirects(response, reverse('dashboard'))


class LoginCsrfTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('giris_kullanici', password='Pass12345!')
        self.client = Client(enforce_csrf_checks=True)
        self.login_url = reverse('login')

    def login_token(self):
        response = self.client.get(self.login_url)
        return response.context['csrf_token']

    def test_double_submit_after_login_redirects_instead_of_error(self):
        old_token = self.login_token()
        data = {'username': 'giris_kullanici', 'password': 'Pass12345!', 'csrfmiddlewaretoken': old_token}
        self.assertRedirects(self.client.post(self.login_url, data), reverse('dashboard'), fetch_redirect_response=False)
        # İkinci basış: giriş anahtarı yenilendiği için eski anahtarla gelir.
        response = self.client.post(self.login_url, data)
        self.assertRedirects(response, reverse('dashboard'), fetch_redirect_response=False)

    def test_double_submit_keeps_safe_next_url(self):
        old_token = self.login_token()
        url = self.login_url + '?next=' + reverse('customer_list')
        data = {'username': 'giris_kullanici', 'password': 'Pass12345!', 'csrfmiddlewaretoken': old_token}
        self.client.post(url, data)
        self.assertRedirects(self.client.post(url, data), reverse('customer_list'), fetch_redirect_response=False)
        evil = self.login_url + '?next=https://example.com/'
        self.assertRedirects(self.client.post(evil, data), reverse('dashboard'), fetch_redirect_response=False)

    def test_stale_login_form_shows_friendly_message(self):
        self.login_token()
        response = self.client.post(self.login_url, {'username': 'giris_kullanici', 'password': 'x', 'csrfmiddlewaretoken': 'eski'})
        self.assertEqual(response.status_code, 403)
        self.assertContains(response, 'güvenlik süresi doldu', status_code=403)
        self.assertContains(response, 'value="giris_kullanici"', status_code=403)
        self.assertContains(response, 'csrfmiddlewaretoken', status_code=403)

    def test_logged_in_user_skips_login_page(self):
        self.client.force_login(self.user)
        self.assertRedirects(self.client.get(self.login_url), reverse('dashboard'), fetch_redirect_response=False)

    def test_stale_form_elsewhere_shows_friendly_page(self):
        self.client.force_login(self.user)
        response = self.client.post(reverse('logout'), {'csrfmiddlewaretoken': 'eski'})
        self.assertContains(response, 'Sayfa yenilenmeli', status_code=403)

    def test_login_page_is_mobile_friendly(self):
        response = self.client.get(self.login_url)
        self.assertContains(response, 'data-submit-once')
        self.assertContains(response, 'data-password-toggle')
        self.assertContains(response, 'autocapitalize="none"')


class BadQueryParameterTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user('param_admin', password='Pass12345!')
        self.admin.profile.role = UserProfile.Role.ADMIN
        self.admin.profile.save()
        self.customer = Customer.objects.create(name='Parametre Müşteri')
        self.client.force_login(self.admin)

    def test_non_numeric_customer_on_order_form_is_not_found(self):
        for value in ('abc', '1.5', 'undefined'):
            self.assertEqual(self.client.get(reverse('order_create') + f'?customer={value}').status_code, 404)

    def test_invalid_history_dates_are_ignored(self):
        url = reverse('customer_history', args=[self.customer.pk])
        self.assertEqual(self.client.get(url + '?date_from=abc&date_to=2026-13-45').status_code, 200)
        self.assertEqual(self.client.get(url + '?date_from=2026-01-01&date_to=2026-12-31').status_code, 200)

    def test_unknown_template_kind_is_not_found(self):
        self.assertEqual(self.client.get(reverse('download_template', args=['unknown'])).status_code, 404)
        self.assertEqual(self.client.get(reverse('download_template', args=['products'])).status_code, 200)


class UploadNameTests(TestCase):
    def test_uploaded_file_names_are_unguessable(self):
        from .models import delivery_proof_upload_to, invoice_upload_to
        first = invoice_upload_to(None, 'Taranan Fatura.PDF')
        self.assertRegex(first, r'^invoices/\d{4}/\d{2}/taranan-fatura-[0-9a-f]{16}\.pdf$')
        self.assertNotEqual(first, invoice_upload_to(None, 'Taranan Fatura.PDF'))
        self.assertRegex(delivery_proof_upload_to(None, 'şğı.jpg'), r'^delivery_proofs/\d{4}/\d{2}/[a-z0-9-]*-?[0-9a-f]{16}\.jpg$')


class SalespersonLabelTests(TestCase):
    def test_salesperson_choices_show_full_name(self):
        user = User.objects.create_user('etiket_satis', first_name='Berkan', last_name='Miralay')
        user.profile.role = UserProfile.Role.SALES
        user.profile.save()
        nameless = User.objects.create_user('isimsiz_satis')
        nameless.profile.role = UserProfile.Role.SALES
        nameless.profile.save()
        labels = [label for _, label in CustomerForm().fields['assigned_salesperson'].choices]
        self.assertIn('Berkan Miralay', labels)
        self.assertIn('isimsiz_satis', labels)
